import uuid
from lxml import etree
from pptx import Presentation

PATH = r"C:\Users\haals\AppData\Roaming\Claude\scratch-workspaces\8ad54198-bc4a-404a-ac38-ba080c4dbefd\40c8920e-8774-4a4f-aa7b-40d3444157cf\scratch-2026-09-10-eb756b\pptx_work\deck_filled.pptx"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
SECTION_EXT_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"


def qn(ns, tag):
    return f"{{{ns}}}{tag}"


prs = Presentation(PATH)
pres_el = prs.slides._sldIdLst.getparent()
slide_ids = [sld.get("id") for sld in prs.slides._sldIdLst]
assert len(slide_ids) == 26, len(slide_ids)

# (section name, how many of the next slides belong to it)
sections = [
    ("Introduction", 3),
    ("The Idea", 2),
    ("The Problems", 5),
    ("Solutions", 4),
    ("The Wiring", 8),
    ("Goals", 3),
    ("Closing", 1),
]
total = sum(n for _, n in sections)
assert total == 26, total

extlst = pres_el.find(qn(P_NS, "extLst"))
assert extlst is not None

ext = etree.SubElement(extlst, qn(P_NS, "ext"))
ext.set("uri", SECTION_EXT_URI)

section_lst = etree.SubElement(ext, qn(P14_NS, "sectionLst"), nsmap={"p14": P14_NS})

idx = 0
for name, count in sections:
    section = etree.SubElement(section_lst, qn(P14_NS, "section"))
    section.set("name", name)
    section.set("id", "{%s}" % str(uuid.uuid4()).upper())
    sld_id_lst = etree.SubElement(section, qn(P14_NS, "sldIdLst"))
    for sid in slide_ids[idx: idx + count]:
        sld_id = etree.SubElement(sld_id_lst, qn(P14_NS, "sldId"))
        sld_id.set("id", sid)
    idx += count

prs.save(PATH)
print("Sections added:", [s[0] for s in sections])
