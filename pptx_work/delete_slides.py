from pptx import Presentation

PATH = r"C:\Users\haals\AppData\Roaming\Claude\scratch-workspaces\8ad54198-bc4a-404a-ac38-ba080c4dbefd\40c8920e-8774-4a4f-aa7b-40d3444157cf\scratch-2026-09-10-eb756b\pptx_work\deck_filled.pptx"

prs = Presentation(PATH)
xml_slides = prs.slides._sldIdLst
slides = list(xml_slides)

# 1-indexed slide numbers 27..45 -> 0-indexed 26..44
to_remove = slides[26:45]
print("Removing", len(to_remove), "slides (27-45)")
for sld in to_remove:
    rId = sld.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    prs.part.drop_rel(rId)
    xml_slides.remove(sld)

prs.save(PATH)
print("Saved. New slide count:", len(prs.slides))
