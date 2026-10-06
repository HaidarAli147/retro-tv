from pptx import Presentation
from pptx.dml.color import RGBColor

PATH = r"C:\Users\haals\AppData\Roaming\Claude\scratch-workspaces\8ad54198-bc4a-404a-ac38-ba080c4dbefd\40c8920e-8774-4a4f-aa7b-40d3444157cf\scratch-2026-09-10-eb756b\pptx_work\deck_filled.pptx"

prs = Presentation(PATH)
slides = list(prs.slides)


def by_id(slide, shape_id):
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    raise KeyError(f"shape id {shape_id} not found")


def set_text(slide, shape_id, new_text):
    shape = by_id(slide, shape_id)
    tf = shape.text_frame
    first_para = tf.paragraphs[0]
    if not first_para.runs:
        first_para.add_run()
    first_para.runs[0].text = new_text
    for r in first_para.runs[1:]:
        r.text = ""
    for p in tf.paragraphs[1:]:
        for r in p.runs:
            r.text = ""


def set_cell(shape, r, c, text):
    shape.table.cell(r, c).text = text


# ---------------------------------------------------------------- Slide 7
# "MOST COMMON CAUSES" was clipped to "MOST / COMMON / SES"
s = slides[6]
set_text(s, 1452, "CAUSES")

# ---------------------------------------------------------------- Slide 9
# Table rows were overlapping - text far too long for the row height
s = slides[8]
t1 = by_id(s, 1635)
set_cell(t1, 1, 1, "Wider than screws")
set_cell(t1, 2, 1, "Size unconfirmed")
set_cell(t1, 3, 1, "Grabbed by desktop")
t2 = by_id(s, 1634)
set_cell(t2, 1, 1, "Wrong HDMI port")
set_cell(t2, 2, 1, "Changed app + system")
set_cell(t2, 3, 1, "Connector latch broke")

# ---------------------------------------------------------------- Slide 10
# Headers/descriptions wrapped so much they overlapped the card below
s = slides[9]
set_text(s, 1783, "mpv was sending sound out the wrong HDMI port")
set_text(s, 1786, "VOLUME CONTROL")
set_text(s, 1785, "The knob changed app and system volume together")
set_text(s, 1788, "SCREEN CONNECTOR")
set_text(s, 1787, "The connector latch snapped during repair")
set_text(s, 1790, "HOLE VS SCREWS")
set_text(s, 1789, "Joystick hole ended up wider than its screws")

# ---------------------------------------------------------------- Slide 12
# "PRIORITY FIX" broke mid-word; headers wrapped into the body text below
s = slides[11]
set_text(s, 1930, "TOP FIX")
set_text(s, 1932, "VOLUME")
set_text(s, 1936, "Volume knob now calls wpctl to change the speaker")
set_text(s, 1933, "AUDIO DEVICE")
set_text(s, 1931, "mpv now points straight at the correct ALSA device")
set_text(s, 1934, "KEY INPUT")
set_text(s, 1935, "Checks the keysym and the raw hardware keycode both")

# ---------------------------------------------------------------- Slide 13
# Long headers were bleeding into the connector lines
s = slides[12]
set_text(s, 1944, "wpctl status to check audio routing")
set_text(s, 1945, "key_debug.py to print real keycodes")
set_text(s, 1946, "Flashlight test for backlight vs video")
set_text(s, 1942, "iptv-org database to verify channels")
set_text(s, 1947, "BFS check for keyboard reachability")
set_text(s, 1943, "Round-trip tests before sending installer")

# ---------------------------------------------------------------- Slide 14
# Body text ran off the bottom of the slide
s = slides[13]
set_text(s, 2038, "Software problems are fixed. Screen repair is in progress.")

# ---------------------------------------------------------------- Slide 19
# "CONTROLLER" broke mid-word onto its own line
s = slides[18]
set_text(s, 2256, "FINISHED BUILD")

# ---------------------------------------------------------------- Slide 20
# This run had no explicit color and rendered near-invisible (dark on dark)
s = slides[19]
shape = by_id(s, 2353)
run = shape.text_frame.paragraphs[0].runs[0]
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# ---------------------------------------------------------------- Slide 22
# Step descriptions overlapped the zig-zag timeline graphic
s = slides[21]
set_text(s, 2405, "Wire inputs to the Arduino")
set_text(s, 2407, "Flash sketch, test serial output")
set_text(s, 2401, "Bridge turns serial into key events")
set_text(s, 2403, "Wire board and screen into case")

prs.save(PATH)
print("Saved fixes to", PATH)
